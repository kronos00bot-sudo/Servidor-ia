#!/usr/bin/env python3
import os, json, requests
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).parent.parent.parent / ".env")

PROJECT_DIR    = Path(os.getenv("PROJECT_DIR"))
DGX_OLLAMA_URL = os.getenv("DGX_OLLAMA_URL", "http://100.64.129.87:11434")
DGX_LLM_MODEL  = os.getenv("DGX_LLM_MODEL", "nemotron-3-super:120b")
DOCS_DIR       = PROJECT_DIR / "data" / "documents"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

PROMPT_SUMMARY = "You are analyzing a document from a WhatsApp conversation. Summarize the key content: main topics, important data, conclusions. Be concise but complete. Answer in English."

def extract_text(filepath):
    ext = filepath.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                return chr(10).join(p.extract_text() or "" for p in pdf.pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(filepath)
            return chr(10).join(p.text for p in doc.paragraphs)
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append("	".join(str(c) if c is not None else "" for c in row))
            return chr(10).join(rows)
        elif ext in (".txt", ".csv"):
            return filepath.read_text(encoding="utf-8", errors="replace")
        else:
            return ""
    except Exception as e:
        return "[Error extrayendo texto: " + str(e) + "]"

def summarize_with_llm(text):
    if not text.strip():
        return "[Documento vacio o sin texto extraible]"
    truncated = text[:8000] + ("..." if len(text) > 8000 else "")
    resp = requests.post(
        DGX_OLLAMA_URL + "/api/generate",
        json={
            "model": DGX_LLM_MODEL,
            "prompt": PROMPT_SUMMARY + chr(10) + chr(10) + "Document content:" + chr(10) + truncated,
            "stream": False
        },
        timeout=180
    )
    resp.raise_for_status()
    return resp.json().get("response", "").strip()

def process_document(filepath, msg_id):
    input_path = Path(filepath)
    result_path = DOCS_DIR / (str(msg_id).zfill(4) + "_" + input_path.stem + ".json")
    if result_path.exists():
        with open(result_path) as f:
            return json.load(f)
    if not input_path.exists():
        return {"id": msg_id, "text": None, "error": "Archivo no encontrado: " + filepath}
    try:
        print("  Extrayendo texto: " + input_path.name)
        raw_text = extract_text(input_path)
        print("  Resumiendo con DGX (" + str(len(raw_text)) + " chars)...")
        summary = summarize_with_llm(raw_text)
        result = {
            "id": msg_id,
            "original": str(filepath),
            "raw_chars": len(raw_text),
            "text": summary,
            "error": None
        }
    except Exception as e:
        result = {"id": msg_id, "original": str(filepath), "text": None, "error": str(e)}
    with open(result_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    return result

def process_documents(messages):
    to_process = [m for m in messages if m["type"] == "document" and m.get("filepath")]
    if not to_process:
        print("No hay documentos para procesar.")
        return messages
    print("Procesando " + str(len(to_process)) + " documentos con DGX...")
    for msg in to_process:
        print("[" + str(msg["id"]).zfill(3) + "] " + msg["sender"] + " - " + msg["filename"])
        result = process_document(msg["filepath"], msg["id"])
        msg["result"] = result.get("text")
        msg["processed"] = True
        if result.get("error"):
            print("  ERROR: " + result["error"])
        else:
            print("  OK: " + result["text"][:80] + "...")
    return messages

if __name__ == "__main__":
    import sys
    parsed_path = PROJECT_DIR / "data" / "conversations" / "parsed.json"
    if not parsed_path.exists():
        print("Ejecuta primero parser.py")
        sys.exit(1)
    with open(parsed_path, encoding="utf-8") as f:
        messages = json.load(f)
    messages = process_documents(messages)
    with open(parsed_path, "w", encoding="utf-8") as f:
        json.dump(messages, f, ensure_ascii=False, indent=2)
    print("Procesamiento de documentos completado.")
