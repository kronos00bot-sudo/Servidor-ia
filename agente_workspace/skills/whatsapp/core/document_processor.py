"""Document processing for WhatsApp skill."""

import json
from pathlib import Path
from skills.whatsapp.routing.remote_client import RemoteClient
from skills.whatsapp.utils.config import Config

CONFIG = Config()
REMOTE = RemoteClient()
DOCS_DIR = CONFIG.PROJECT_DIR / 'data' / 'documents'
DOCS_DIR.mkdir(parents=True, exist_ok=True)
PROMPT_SUMMARY = "You are analyzing a document from a WhatsApp conversation. Summarize the key content: main topics, important data, conclusions. Be concise but complete. Answer in English."


def extract_text(filepath: Path) -> str:
    ext = filepath.suffix.lower()
    try:
        if ext == '.pdf':
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                return '\n'.join(p.extract_text() or '' for p in pdf.pages)
        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            return '\n'.join(p.text for p in doc.paragraphs)
        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append('\t'.join(str(c) if c is not None else '' for c in row))
            return '\n'.join(rows)
        elif ext in ('.txt', '.csv'):
            return filepath.read_text(encoding='utf-8', errors='replace')
        else:
            return ''
    except Exception as e:
        return f'[Error extracting text: {e}]'


def summarize_with_llm(text: str, router) -> str:
    if not text.strip():
        return '[Documento vacío o sin texto extraíble]'
    truncated = text[:8000] + ('...' if len(text) > 8000 else '')
    route = router.route_document({"kind": "document"})
    prompt = PROMPT_SUMMARY + '\n\nDocument content:\n' + truncated
    try:
        raw = REMOTE.generate(route.host, route.model, prompt, timeout=route.timeout)
    except Exception:
        fallback = router.route_fallback("document")
        raw = REMOTE.generate(fallback.host, fallback.model, prompt, timeout=fallback.timeout)
    return raw.get('response', '').strip()


def process_document(filepath: str, msg_id: int, router) -> dict:
    input_path = Path(filepath)
    result_path = DOCS_DIR / f"{msg_id:04d}_{input_path.stem}.json"
    if result_path.exists():
        with open(result_path, encoding='utf-8') as f:
            return json.load(f)
    if not input_path.exists():
        return {'id': msg_id, 'text': None, 'error': f'File not found: {filepath}'}
    try:
        raw_text = extract_text(input_path)
        summary_text = summarize_with_llm(raw_text, router)
        result = {
            'id': msg_id,
            'original': str(filepath),
            'raw_chars': len(raw_text),
            'text': summary_text,
            'error': None
        }
    except Exception as e:
        result = {'id': msg_id, 'original': str(filepath), 'text': None, 'error': str(e)}
    with open(result_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    return result


def process_documents(messages: list, router) -> list:
    tasks = [m for m in messages if m['type'] == 'document' and m.get('filepath')]
    if not tasks:
        print('No documents to process.')
        return messages

    print(f'Processing {len(tasks)} documents with DGX...')
    for msg in tasks:
        print(f"[{msg['id']:03d}] {msg['sender']} - {msg['filename']}")
        result = process_document(msg['filepath'], msg['id'], router)
        msg['result'] = result.get('text')
        msg['processed'] = True
        if result.get('error'):
            print(f"  ERROR: {result['error']}")
        else:
            print(f"  OK: {msg['result'][:80]}...")
    return messages
